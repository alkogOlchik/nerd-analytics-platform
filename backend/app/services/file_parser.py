"""Parse uploaded files to plain text."""

import csv
import io
from pathlib import PurePosixPath

MAX_FILE_TEXT_CHARS: int = 50_000

_SUPPORTED = {".pdf", ".docx", ".doc", ".txt", ".md", ".rst", ".markdown", ".xlsx", ".xls", ".csv"}


def _ext(filename: str) -> str:
    return PurePosixPath(filename).suffix.lower()


def _parse_pdf(data: bytes) -> str:
    try:
        from pypdf import PdfReader
    except ImportError as e:
        raise ValueError("pypdf not installed") from e
    reader = PdfReader(io.BytesIO(data))
    parts = [page.extract_text() or "" for page in reader.pages]
    return "\n".join(parts)


def _parse_docx(data: bytes) -> str:
    try:
        from docx import Document
    except ImportError as e:
        raise ValueError("python-docx not installed") from e
    doc = Document(io.BytesIO(data))
    return "\n".join(p.text for p in doc.paragraphs)


def _parse_xlsx(data: bytes) -> str:
    try:
        import openpyxl
    except ImportError as e:
        raise ValueError("openpyxl not installed") from e
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    parts: list[str] = []
    for sheet in wb.worksheets:
        parts.append(f"[{sheet.title}]")
        for row in sheet.iter_rows(values_only=True):
            parts.append("\t".join("" if v is None else str(v) for v in row))
    return "\n".join(parts)


def _parse_csv(data: bytes) -> str:
    text = data.decode("utf-8", errors="replace")
    reader = csv.reader(io.StringIO(text))
    return "\n".join("\t".join(row) for row in reader)


def parse_to_text(data: bytes, filename: str) -> str:
    ext = _ext(filename)
    if ext not in _SUPPORTED:
        raise ValueError(f"Unsupported file type: {ext}")
    try:
        if ext == ".pdf":
            text = _parse_pdf(data)
        elif ext in (".docx", ".doc"):
            text = _parse_docx(data)
        elif ext in (".txt", ".md", ".rst", ".markdown"):
            text = data.decode("utf-8", errors="replace")
        elif ext in (".xlsx", ".xls"):
            text = _parse_xlsx(data)
        elif ext == ".csv":
            text = _parse_csv(data)
        else:
            text = ""
    except ValueError:
        raise
    except Exception as e:
        raise ValueError(f"Failed to parse {filename}: {e}") from e

    return text[:MAX_FILE_TEXT_CHARS]
